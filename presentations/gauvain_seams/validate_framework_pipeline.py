#!/usr/bin/env python3
"""Validate the standalone editable framework-pipeline PowerPoint deliverable."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "Gauvain_Framework_Pipeline_Current_Editable.pptx"
PREVIEW = ROOT / "Gauvain_Framework_Pipeline_Current_Editable_preview.png"
CONTACT = ROOT / "Gauvain_Framework_Pipeline_Current_Editable_contact_sheet.png"

EXPECTED_FLOW_OBJECTS = {
    "Paper input pinhole video frame",
    "Paper input video to frame extraction arrow",
    "Paper frame extraction panel",
    "Paper frame extraction to COLMAP arrow",
    "Paper COLMAP panel",
    "Paper COLMAP to posed frames arrow",
    "Paper posed real frames panel",
    "Paper posed frames to initial scene arrow",
    "Paper initial 3DGRUT panel",
    "Paper initial scene to shared scene arrow",
    "Paper shared 3D Gaussian scene panel",
    "Paper shared scene to multi-direction renders arrow",
    "Paper spherical snake panel",
    "Paper spherical snake to ArtiFixer arrow",
    "Paper ArtiFixer panel",
    "Paper ArtiFixer to restored observations arrow",
    "Paper restored observations panel",
    "Paper restored observations to distillation arrow",
    "Paper ArtiFixer3D distillation panel",
    "Paper distillation feedback to shared scene arrow",
    "Paper real anchor branch to ArtiFixer arrow",
    "Paper real anchor branch to distillation arrow",
    "Paper refined scene to synchronized rendering arrow",
    "Paper synchronized panoramic rendering panel",
    "Paper synchronized rendering to final ArtiFixer3D plus arrow",
    "Paper final ArtiFixer3D plus panel",
    "Paper final ArtiFixer3D plus to spherical fusion arrow",
    "Paper spherical fusion ERP panel",
    "Paper ERP projection to output arrow",
    "Paper output ERP video frame",
}

REQUIRED_TEXT = (
    "Pinhole input video",
    "Frame\nextraction",
    "COLMAP",
    "Posed real frames",
    "Initial 3DGRUT /\n3D Gaussian Scene",
    "Shared 3D\nGaussian Scene",
    "World-locked overlapping\npinhole renders",
    "spherical snake: stable directions + smooth turns",
    "ArtiFixer",
    "video-to-video",
    "one common temporal stream",
    "Restored\npseudo-\nobservations",
    "Geometry-constrained\nArtiFixer3D distillation",
    "geometry preserved · appearance refined",
    "real-image anchors",
    "Synchronized\npanoramic rendering",
    "Final\nArtiFixer3D+",
    "Spherical fusion / ERP projection",
    "Equirectangular 360°\nvideo (ERP)",
    "No native spherical topology",
)

FORBIDDEN_TEXT_PATTERNS = (
    re.compile(r"\bcubemap\b", re.I),
    re.compile(r"cube\s*(?:-|→|to)\s*ERP", re.I),
    re.compile(r"\b154\b|\b210\b|\b14 directions\b|\b15 fps\b", re.I),
    re.compile(r"bearlake|room tour", re.I),
)


def validate() -> None:
    assert PPTX.is_file(), f"missing {PPTX}"
    first = Presentation(PPTX)
    assert len(first.slides) == 1, f"expected one slide, found {len(first.slides)}"
    slide = first.slides[0]
    assert abs(first.slide_width / first.slide_height - 16 / 9) < 1e-6, "slide is not 16:9"

    names = [shape.name for shape in slide.shapes]
    assert len(names) == len(set(names)), "shape names are not unique"
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.GROUP for shape in slide.shapes), "group found"
    missing = sorted(EXPECTED_FLOW_OBJECTS - set(names))
    assert not missing, f"missing flow objects: {missing}"

    slide_texts = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
    full_text = "\n".join(slide_texts)
    for text in REQUIRED_TEXT:
        assert text in full_text, f"missing required text: {text!r}"
    for pattern in FORBIDDEN_TEXT_PATTERNS:
        assert not pattern.search(full_text), f"forbidden experimental/cubemap text: {pattern.pattern}"

    picture_shapes = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(picture_shapes) == 2, f"expected two pictures, found {len(picture_shapes)}"
    for picture in picture_shapes:
        assert picture.width > 0 and picture.height > 0, f"invalid image geometry: {picture.name}"
        assert all(
            0.0 <= value <= 1.0
            for value in (picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom)
        ), f"invalid crop fraction: {picture.name}"

    slide_w, slide_h = first.slide_width, first.slide_height
    overflow = []
    for shape in slide.shapes:
        if shape.left < 0 or shape.top < 0 or shape.left + shape.width > slide_w or shape.top + shape.height > slide_h:
            overflow.append(shape.name)
    assert not overflow, f"objects outside slide bounds: {overflow}"

    with zipfile.ZipFile(PPTX) as archive:
        media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
        assert len(media) == 2, f"expected two embedded media files, found {media}"
        for member in media:
            with Image.open(io.BytesIO(archive.read(member))) as image:
                assert image.width > 0 and image.height > 0

    # Reopen independently to verify serialization and package relationships.
    second = Presentation(PPTX)
    assert len(second.slides) == 1
    assert len(second.slides[0].shapes) == len(slide.shapes)

    for image_path in (PREVIEW, CONTACT):
        assert image_path.is_file(), f"missing {image_path}"
        with Image.open(image_path) as image:
            assert image.size == (1600, 900), f"unexpected render size for {image_path}: {image.size}"
            assert image.width / image.height == 16 / 9

    print(
        "PASS",
        {
            "slides": len(first.slides),
            "independent_shapes": len(slide.shapes),
            "groups": 0,
            "pictures": len(picture_shapes),
            "flow_objects": len(EXPECTED_FLOW_OBJECTS),
            "preview": str(PREVIEW),
            "contact_sheet": str(CONTACT),
        },
    )


if __name__ == "__main__":
    validate()
